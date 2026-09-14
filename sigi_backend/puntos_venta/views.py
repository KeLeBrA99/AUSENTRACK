"""
AUSENTRACK - Vistas de Puntos de Venta
"""
import re
from io import BytesIO

from django.http import HttpResponse
from django.db.models import Q
from rest_framework import generics, status
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated

from authentication.permissions import EsAdmin, EsTalentoHumano
from authentication.views import registrar_auditoria
from .models import PuntoVenta
from .serializers import PuntoVentaSerializer

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    OPENPYXL_DISPONIBLE = True
except ImportError:
    OPENPYXL_DISPONIBLE = False


def _normalizar(texto):
    """
    Normaliza un nombre de punto de venta para poder comparar 'Poke 2',
    'Poke2' y 'POKE 2' como el mismo punto, sin depender de que el texto
    libre de 'area' en Colaborador este escrito siempre igual.
    """
    if not texto:
        return ''
    return re.sub(r'[^a-z0-9]', '', texto.lower())


# ── CRUD ─────────────────────────────────────────────────────────────────────

class ListarPuntosVentaView(generics.ListAPIView):
    """
    GET /api/puntos-venta/
    Filtros: ?activo=true&search=
    """
    serializer_class   = PuntoVentaSerializer
    permission_classes = [IsAuthenticated, EsTalentoHumano]

    def get_queryset(self):
        qs = PuntoVenta.objects.all()
        activo = self.request.query_params.get('activo')
        search = self.request.query_params.get('search')
        if activo is not None:
            qs = qs.filter(activo=activo.lower() == 'true')
        if search:
            qs = qs.filter(Q(nombre__icontains=search))
        return qs


class CrearPuntoVentaView(generics.CreateAPIView):
    serializer_class   = PuntoVentaSerializer
    permission_classes = [IsAuthenticated, EsTalentoHumano]

    def perform_create(self, serializer):
        pv = serializer.save()
        registrar_auditoria(
            self.request.user, 'CREATE', 'punto_venta', pv.id,
            detalle=f'Punto de venta creado: {pv.nombre}', request=self.request,
        )


class DetallePuntoVentaView(generics.RetrieveUpdateDestroyAPIView):
    queryset           = PuntoVenta.objects.all()
    serializer_class   = PuntoVentaSerializer
    permission_classes = [IsAuthenticated, EsTalentoHumano]

    def get_permissions(self):
        if self.request.method == 'DELETE':
            return [IsAuthenticated(), EsAdmin()]
        return [IsAuthenticated(), EsTalentoHumano()]

    def perform_update(self, serializer):
        pv = serializer.save()
        registrar_auditoria(
            self.request.user, 'UPDATE', 'punto_venta', pv.id,
            detalle=f'Punto de venta actualizado: {pv.nombre}', request=self.request,
        )

    def perform_destroy(self, instance):
        registrar_auditoria(
            self.request.user, 'DELETE', 'punto_venta', instance.id,
            detalle=f'Punto de venta desactivado: {instance.nombre}', request=self.request,
        )
        instance.activo = False
        instance.save(update_fields=['activo'])


# ── COBERTURA (personal requerido vs. real) ──────────────────────────────────

class CoberturaPuntosVentaView(APIView):
    """
    GET /api/puntos-venta/cobertura/
    Compara personal_requerido de cada punto contra los colaboradores
    activos cuya 'area' coincide (comparacion flexible: ignora mayusculas,
    espacios y puntuacion, para que 'Poke 2' y 'Poke2' cuenten igual).
    """
    permission_classes = [IsAuthenticated, EsTalentoHumano]

    def get(self, request):
        from colaboradores.models import Colaborador

        puntos = PuntoVenta.objects.filter(activo=True)
        colaboradores_activos = Colaborador.objects.filter(activo=True)

        # Conteo por FK real (fuente de verdad) + fallback por texto para
        # colaboradores antiguos que aun no tienen punto_venta asignado.
        conteo_por_punto_id = {}
        conteo_por_area_texto = {}
        for c in colaboradores_activos:
            if c.punto_venta_id:
                conteo_por_punto_id[c.punto_venta_id] = conteo_por_punto_id.get(c.punto_venta_id, 0) + 1
            else:
                key = _normalizar(c.area)
                if key:
                    conteo_por_area_texto[key] = conteo_por_area_texto.get(key, 0) + 1

        nombres_normalizados = {_normalizar(p.nombre) for p in puntos}

        resultado = []
        total_requerido = 0
        total_actual = 0
        con_falta = 0
        con_exceso = 0

        for p in puntos:
            key = _normalizar(p.nombre)
            actual = conteo_por_punto_id.get(p.id, 0) + conteo_por_area_texto.get(key, 0)
            diferencia = actual - p.personal_requerido
            if diferencia < 0:
                estado = 'FALTA'
                con_falta += 1
            elif diferencia > 0:
                estado = 'EXCESO'
                con_exceso += 1
            else:
                estado = 'COMPLETO'

            resultado.append({
                'id': p.id, 'nombre': p.nombre,
                'personal_requerido': p.personal_requerido,
                'personal_actual': actual,
                'diferencia': diferencia,
                'estado': estado,
            })
            total_requerido += p.personal_requerido
            total_actual += actual

        resultado.sort(key=lambda x: x['diferencia'])

        # Colaboradores sin punto_venta asignado Y cuya 'area' de texto tampoco
        # coincide con ningun punto de venta registrado (datos realmente huerfanos)
        sin_punto_valido = []
        for c in colaboradores_activos:
            if c.punto_venta_id:
                continue
            key = _normalizar(c.area)
            if not key or key not in nombres_normalizados:
                sin_punto_valido.append({'id_colaborador': c.id_colaborador, 'nombre': c.nombre, 'area': c.area or '(vacio)'})

        return Response({
            'kpis': {
                'total_puntos': puntos.count(),
                'personal_requerido_total': total_requerido,
                'personal_actual_total': total_actual,
                'puntos_con_falta': con_falta,
                'puntos_con_exceso': con_exceso,
            },
            'puntos': resultado,
            'colaboradores_sin_punto_valido': sin_punto_valido,
        })


# ── IMPORTAR / PLANTILLA ──────────────────────────────────────────────────────

class DescargarPlantillaPuntosVentaView(APIView):
    permission_classes = [IsAuthenticated, EsTalentoHumano]

    def get(self, request):
        if not OPENPYXL_DISPONIBLE:
            return Response({'detail': 'openpyxl no esta instalado.'}, status=500)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Plantilla'

        ws.merge_cells('A1:E1')
        ws['A1'] = 'AUSENTRACK - Plantilla de Puntos de Venta'
        ws['A1'].font = Font(bold=True, size=13, color='4F8B5B')

        ws.merge_cells('A2:E2')
        ws['A2'] = 'Complete desde la fila 6. No modifique los nombres tecnicos de la fila 4.'
        ws['A2'].font = Font(italic=True, size=10, color='6B6F76')

        columnas = [
            ('nombre',             'Nombre *',              True,  22),
            ('personal_requerido', 'Personal Requerido *',  True,  18),
            ('direccion',          'Direccion',             False, 28),
            ('telefono',           'Telefono',              False, 16),
            ('notas',              'Notas',                 False, 30),
        ]
        header_fill = PatternFill(start_color='4F8B5B', end_color='4F8B5B', fill_type='solid')
        header_font = Font(color='FFFFFF', bold=True)
        visual_fill = PatternFill(start_color='EAF3EC', end_color='EAF3EC', fill_type='solid')

        for i, (tecnico, visual, obligatorio, ancho) in enumerate(columnas, start=1):
            c1 = ws.cell(row=4, column=i, value=tecnico)
            c1.fill = header_fill
            c1.font = header_font
            c2 = ws.cell(row=5, column=i, value=visual)
            c2.fill = visual_fill
            c2.font = Font(bold=obligatorio, size=10)
            ws.column_dimensions[c2.column_letter].width = ancho

        ejemplos = [
            ('Poke 1', 5, '', '', ''),
            ('Poke 2', 6, '', '', ''),
        ]
        for r, fila in enumerate(ejemplos, start=6):
            for c, valor in enumerate(fila, start=1):
                ws.cell(row=r, column=c, value=valor)

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        response = HttpResponse(
            buffer.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = 'attachment; filename="plantilla_puntos_venta.xlsx"'
        return response


class ImportarPuntosVentaView(APIView):
    """
    POST /api/puntos-venta/importar/
    Si el nombre ya existe (comparacion flexible), ACTUALIZA el personal_requerido
    en vez de duplicar el punto de venta.
    """
    permission_classes = [IsAuthenticated, EsTalentoHumano]

    def post(self, request):
        archivo = request.FILES.get('archivo')
        if not archivo:
            return Response({'error': 'No se recibio ningun archivo.'}, status=400)
        if not OPENPYXL_DISPONIBLE:
            return Response({'detail': 'openpyxl no esta instalado.'}, status=500)

        try:
            wb = openpyxl.load_workbook(archivo, data_only=True)
        except Exception:
            return Response({'error': 'No se pudo leer el archivo. Verifica que sea un .xlsx valido.'}, status=400)

        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))

        encabezados = None
        header_idx = None
        for idx, row in enumerate(rows):
            valores = [str(c).strip().lower() if c else '' for c in row]
            if 'nombre' in valores and 'personal_requerido' in valores:
                encabezados = valores
                header_idx = idx
                break

        if encabezados is None:
            return Response({
                'error': 'No se encontraron los encabezados. El archivo debe tener una fila con, al menos, las columnas: nombre, personal_requerido.'
            }, status=400)

        def col(nombre):
            try:
                return encabezados.index(nombre)
            except ValueError:
                return None

        i_nombre = col('nombre')
        i_personal = col('personal_requerido')
        i_direccion = col('direccion')
        i_telefono = col('telefono')
        i_notas = col('notas')

        filas_datos = rows[header_idx + 1:]

        def parece_fila_de_etiqueta(fila):
            valor = fila[i_nombre] if i_nombre is not None and i_nombre < len(fila) else None
            if valor is None:
                return True
            texto = str(valor).strip()
            return texto == '' or texto.endswith('*')

        while filas_datos and parece_fila_de_etiqueta(filas_datos[0]):
            filas_datos = filas_datos[1:]

        creados = 0
        actualizados = 0
        errores = []

        existentes = {_normalizar(p.nombre): p for p in PuntoVenta.objects.all()}

        for num_fila, fila in enumerate(filas_datos, start=header_idx + 2):
            nombre = str(fila[i_nombre]).strip() if i_nombre is not None and fila[i_nombre] else None
            if not nombre:
                continue

            try:
                personal = int(float(fila[i_personal])) if i_personal is not None and fila[i_personal] not in (None, '') else 1
            except (TypeError, ValueError):
                errores.append(f'Fila {num_fila}: personal_requerido invalido, se uso 1 por defecto.')
                personal = 1

            key = _normalizar(nombre)
            existente = existentes.get(key)
            if existente:
                existente.personal_requerido = personal
                if i_direccion is not None and fila[i_direccion]:
                    existente.direccion = str(fila[i_direccion]).strip()
                if i_telefono is not None and fila[i_telefono]:
                    existente.telefono = str(fila[i_telefono]).strip()
                if i_notas is not None and fila[i_notas]:
                    existente.notas = str(fila[i_notas]).strip()
                existente.activo = True
                existente.save()
                actualizados += 1
            else:
                nuevo = PuntoVenta.objects.create(
                    nombre=nombre,
                    personal_requerido=personal,
                    direccion=str(fila[i_direccion]).strip() if i_direccion is not None and fila[i_direccion] else '',
                    telefono=str(fila[i_telefono]).strip() if i_telefono is not None and fila[i_telefono] else '',
                    notas=str(fila[i_notas]).strip() if i_notas is not None and fila[i_notas] else '',
                )
                existentes[key] = nuevo
                creados += 1

        registrar_auditoria(
            request.user, 'CREATE', 'punto_venta_import', None,
            detalle=f'Importacion: {creados} creados, {actualizados} actualizados', request=request,
        )

        return Response({'creados': creados, 'actualizados': actualizados, 'errores': errores}, status=200)
